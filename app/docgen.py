"""Relleno de la plantilla ACTA DE EJEMPLO.docx con los datos de una acta.

La plantilla (`templates_docx/acta_template.docx`) es una copia exacta del acta de ejemplo
entregada por el coordinador. Su estructura de tablas fue inspeccionada manualmente:

Tabla 0 (encabezado):
  fila 0: títulos "ACTA DE REUNIÓN No." | "PROCESO(S)/ESTAMENTO" | "LUGAR"
  fila 1: [Acta No. (vacío)] | "Convivencia" | "Coordinación"
  fila 2: títulos "FECHA" | "HORA INICIAL" | "HORA FINAL" | "ELABORADA POR"
  fila 3: valores de la fila anterior

Tabla 1 (cuerpo):
  fila 0: "ASISTENCIA"
  fila 1: títulos "Presentes" | "Cargo/Rol" | "Firma"
  filas 2-7: filas de asistentes (5 con datos de ejemplo + 1 en blanco, clonable)
  fila 8: "ORDEN DEL DÍA" (se deja igual, es una lista fija de 3 puntos)
  fila 9: celda única con la lista del orden del día (se deja igual)
  fila 10: "DESARROLLO"
  fila 11: celda única con los párrafos: Saludo (fijo) / Antecedentes y descripción de la
           situación (reemplazable) / Acuerdos (reemplazable)
  fila 12: fila en blanco de separación
  fila 13: títulos "Descripción" | "Responsable(s)"
  filas 14-16: filas en blanco para acciones concretas (clonables)
  fila 17: "Próxima reunión: " | "HORA: " | "LUGAR: " (se completa agregando texto tras la etiqueta)
"""
import copy
import re
from pathlib import Path

from docx import Document

TEMPLATE_PATH = Path(__file__).resolve().parent / "templates_docx" / "acta_template.docx"

MESES = {
    "enero": "01", "febrero": "02", "marzo": "03", "abril": "04", "mayo": "05", "junio": "06",
    "julio": "07", "agosto": "08", "septiembre": "09", "setiembre": "09", "octubre": "10",
    "noviembre": "11", "diciembre": "12",
}

CARACTERES_INVALIDOS_ARCHIVO = re.compile(r'[<>:"/\\|?*]')


def _formatear_hora(hora: str) -> str:
    """Convierte una hora en formato 24h ('08:00', de <input type=time>) a 12h con am/pm/m,
    replicando la convención del acta de ejemplo ('11:00 am', '12:00 m' para el mediodía)."""
    if not hora:
        return ""
    m = re.match(r"^(\d{1,2}):(\d{2})", hora.strip())
    if not m:
        return hora
    h, minuto = int(m.group(1)), m.group(2)
    if h == 12 and minuto == "00":
        return "12:00 m"
    sufijo = "am" if h < 12 else "pm"
    h12 = h % 12
    if h12 == 0:
        h12 = 12
    return f"{h12}:{minuto} {sufijo}"


def _fecha_para_archivo(fecha_texto: str) -> str:
    """Convierte una fecha en texto ('5 de septiembre de 2026' o '05/09/2026') al formato
    dd.mm.aaaa usado en el nombre de archivo."""
    fecha_texto = (fecha_texto or "").strip()
    m = re.search(r"(\d{1,2})\s+de\s+([a-zA-Záéíóúñ]+)\s+de\s+(\d{4})", fecha_texto, re.IGNORECASE)
    if m:
        dia, mes_txt, anio = m.groups()
        mes = MESES.get(mes_txt.lower(), "00")
        return f"{int(dia):02d}.{mes}.{anio}"
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", fecha_texto)
    if m:
        dia, mes, anio = m.groups()
        return f"{int(dia):02d}.{int(mes):02d}.{anio}"
    return fecha_texto.replace(" ", "_") or "sin-fecha"


def construir_nombre_archivo(acta: dict) -> str:
    """Nombre de archivo: '{numero}-{dd.mm.aaaa} {Nombre del estudiante} {grado}.docx'.
    Si no hay un estudiante con nombre, usa 'Sin Estudiante' en su lugar."""
    fecha_str = _fecha_para_archivo(acta.get("fecha", ""))
    estudiante = next(
        (p for p in acta.get("personas", [])
         if (p.get("tipo") or "").lower() == "estudiante" and (p.get("nombre") or "").strip()),
        None,
    )
    if estudiante:
        nombre = estudiante["nombre"].strip()
        grado_match = re.search(r"(\d+)", estudiante.get("rol_cargo", "") or "")
        grado = grado_match.group(1) if grado_match else ""
        parte_estudiante = f"{nombre} {grado}".strip()
    else:
        parte_estudiante = "Sin Estudiante"

    nombre_archivo = f"{acta.get('numero_acta', '')}-{fecha_str} {parte_estudiante}.docx"
    return CARACTERES_INVALIDOS_ARCHIVO.sub("", nombre_archivo)


def _unique_cells(row):
    """Devuelve las celdas únicas de una fila (deduplicando celdas fusionadas por gridSpan)."""
    seen = []
    for cell in row.cells:
        if not seen or seen[-1]._tc is not cell._tc:
            seen.append(cell)
    return seen


def _set_espaciado_doble(cell):
    for p in cell.paragraphs:
        p.paragraph_format.line_spacing = 2.0


def _set_cell_text(cell, text: str):
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    template_run = first.runs[0] if first.runs else None
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for run in list(first.runs):
        run._r.getparent().remove(run._r)
    new_run = first.add_run(text)
    if template_run is not None:
        new_run.bold = template_run.bold
        new_run.italic = template_run.italic
        if template_run.font.size:
            new_run.font.size = template_run.font.size
        if template_run.font.name:
            new_run.font.name = template_run.font.name


def _append_text_to_cell(cell, extra_text: str):
    first = cell.paragraphs[0]
    template_run = first.runs[-1] if first.runs else None
    new_run = first.add_run(extra_text)
    if template_run is not None:
        new_run.bold = template_run.bold
        new_run.italic = template_run.italic
        if template_run.font.size:
            new_run.font.size = template_run.font.size


def _set_paragraph_multiline(paragraph, text: str):
    template_run = paragraph.runs[0] if paragraph.runs else None
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    lines = text.split("\n")
    new_run = paragraph.add_run(lines[0])
    for line in lines[1:]:
        paragraph.add_run().add_break()
        paragraph.add_run(line)
    if template_run is not None:
        for run in paragraph.runs:
            run.bold = template_run.bold
            run.italic = template_run.italic
            if template_run.font.size:
                run.font.size = template_run.font.size
            if template_run.font.name:
                run.font.name = template_run.font.name


def _eliminar_parrafo(paragraph):
    """Elimina un párrafo completo del documento (a diferencia de solo vaciar su texto)."""
    p = paragraph._p
    p.getparent().remove(p)


def _clone_row_after(table, row_index: int):
    """Clona la fila en `row_index` y la inserta justo después. Devuelve la nueva fila (docx Row)."""
    src_tr = table.rows[row_index]._tr
    new_tr = copy.deepcopy(src_tr)
    src_tr.addnext(new_tr)
    from docx.table import _Row
    return _Row(new_tr, table)


def _find_row_index(table, text_startswith: str, start=0) -> int:
    for i in range(start, len(table.rows)):
        row_text = "".join(c.text for c in _unique_cells(table.rows[i]))
        if row_text.strip().upper().startswith(text_startswith.upper()):
            return i
    raise ValueError(f"No se encontró una fila que empiece con {text_startswith!r}")


def generar_acta_docx(acta: dict, output_path: Path) -> Path:
    """
    acta: dict con las llaves numero_acta, fecha, hora_inicio, hora_fin, lugar, elaborada_por,
          personas (lista de {nombre, rol_cargo}), narrativa_antecedentes, acuerdos_texto,
          acciones (lista de {descripcion, responsable_sugerido}), proxima_reunion (opcional:
          {fecha, hora, lugar}).
    """
    doc = Document(TEMPLATE_PATH)

    # --- Tabla 0: encabezado ---
    t0 = doc.tables[0]
    row1 = _unique_cells(t0.rows[1])
    _set_cell_text(row1[0], str(acta.get("numero_acta", "")))
    # row1[1] = "Convivencia" y row1[2] = "Coordinación" quedan fijos, pero se respeta el lugar si se indicó.
    if acta.get("lugar"):
        _set_cell_text(row1[2], acta["lugar"])

    row3 = _unique_cells(t0.rows[3])
    _set_cell_text(row3[0], acta.get("fecha", ""))
    _set_cell_text(row3[1], _formatear_hora(acta.get("hora_inicio", "")))
    _set_cell_text(row3[2], _formatear_hora(acta.get("hora_fin", "")))
    _set_cell_text(row3[3], acta.get("elaborada_por", "") or "")

    # --- Tabla 1: asistencia ---
    t1 = doc.tables[1]
    personas = acta.get("personas", [])
    primera_fila_asistencia = _find_row_index(t1, "PRESENTES") + 1
    fila_orden_dia = _find_row_index(t1, "ORDEN DEL DÍA")
    filas_disponibles = fila_orden_dia - primera_fila_asistencia

    if len(personas) > filas_disponibles:
        faltantes = len(personas) - filas_disponibles
        ultima_fila_asistencia = fila_orden_dia - 1
        for _ in range(faltantes):
            _clone_row_after(t1, ultima_fila_asistencia)
            ultima_fila_asistencia += 1
        fila_orden_dia = _find_row_index(t1, "ORDEN DEL DÍA")
        filas_disponibles = fila_orden_dia - primera_fila_asistencia

    for i in range(filas_disponibles):
        fila_idx = primera_fila_asistencia + i
        celdas = _unique_cells(t1.rows[fila_idx])
        if i < len(personas):
            persona = personas[i]
            _set_cell_text(celdas[0], persona.get("nombre", ""))
            _set_cell_text(celdas[1], persona.get("rol_cargo", ""))
            _set_cell_text(celdas[2], "")
        else:
            _set_cell_text(celdas[0], "")
            _set_cell_text(celdas[1], "")
            _set_cell_text(celdas[2], "")
        for celda in celdas:
            _set_espaciado_doble(celda)

    # --- Tabla 1: desarrollo (Antecedentes y Acuerdos) ---
    fila_desarrollo = _find_row_index(t1, "DESARROLLO") + 1
    celda_desarrollo = _unique_cells(t1.rows[fila_desarrollo])[0]
    paragraphs = celda_desarrollo.paragraphs

    idx_antecedentes = next(
        i for i, p in enumerate(paragraphs)
        if p.text.strip().lower().startswith("antecedentes y descripci")
    )
    idx_acuerdos = next(
        i for i, p in enumerate(paragraphs)
        if p.text.strip().lower() == "acuerdos"
    )

    narrativa = acta.get("narrativa_antecedentes") or ""
    if idx_antecedentes + 2 < len(paragraphs):
        _set_paragraph_multiline(paragraphs[idx_antecedentes + 2], narrativa)
    # Deja un único párrafo en blanco (dos saltos de línea) entre el contenido y "Acuerdos";
    # el resto de párrafos de relleno de la plantilla de ejemplo se eliminan por completo.
    for j in range(idx_antecedentes + 4, idx_acuerdos):
        _eliminar_parrafo(paragraphs[j])

    acuerdos = acta.get("acuerdos_texto") or ""
    if idx_acuerdos + 2 < len(paragraphs):
        _set_paragraph_multiline(paragraphs[idx_acuerdos + 2], acuerdos)
    # Igual al final: deja como mucho un único párrafo en blanco tras los acuerdos.
    for j in range(idx_acuerdos + 4, len(paragraphs)):
        _eliminar_parrafo(paragraphs[j])

    # --- Tabla 1: acciones (Descripción / Responsable(s)) ---
    fila_titulo_desc = _find_row_index(t1, "DESCRIPCI")
    primera_fila_desc = fila_titulo_desc + 1
    fila_proxima = _find_row_index(t1, "PRÓXIMA REUNIÓN") if _tiene_fila_proxima(t1) else len(t1.rows)
    acciones = acta.get("acciones", [])
    filas_desc_disponibles = fila_proxima - primera_fila_desc

    if len(acciones) > filas_desc_disponibles:
        faltantes = len(acciones) - filas_desc_disponibles
        ultima_fila_desc = fila_proxima - 1
        for _ in range(faltantes):
            _clone_row_after(t1, ultima_fila_desc)
            ultima_fila_desc += 1
        fila_proxima = _find_row_index(t1, "PRÓXIMA REUNIÓN") if _tiene_fila_proxima(t1) else len(t1.rows)
        filas_desc_disponibles = fila_proxima - primera_fila_desc

    for i in range(filas_desc_disponibles):
        fila_idx = primera_fila_desc + i
        celdas = _unique_cells(t1.rows[fila_idx])
        if i < len(acciones):
            accion = acciones[i]
            descripcion = accion.get("descripcion", "")
            categoria = accion.get("categoria")
            articulo = accion.get("articulo")
            if categoria or articulo:
                descripcion = f"{descripcion} ({categoria or ''} {articulo or ''})".strip()
            _set_cell_text(celdas[0], descripcion)
            _set_cell_text(celdas[1], accion.get("responsable_sugerido", ""))
        else:
            _set_cell_text(celdas[0], "")
            _set_cell_text(celdas[1], "")

    # --- Tabla 1: próxima reunión ---
    proxima = acta.get("proxima_reunion")
    if proxima and _tiene_fila_proxima(t1):
        fila_proxima_idx = _find_row_index(t1, "PRÓXIMA REUNIÓN")
        celdas = _unique_cells(t1.rows[fila_proxima_idx])
        if proxima.get("fecha"):
            _append_text_to_cell(celdas[0], proxima["fecha"])
        if proxima.get("hora"):
            _append_text_to_cell(celdas[1], proxima["hora"])
        if proxima.get("lugar"):
            _append_text_to_cell(celdas[2], proxima["lugar"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _tiene_fila_proxima(table) -> bool:
    try:
        _find_row_index(table, "PRÓXIMA REUNIÓN")
        return True
    except ValueError:
        return False
